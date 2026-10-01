"""Policy wrappers exposing one `openpyxl` workbook to the Monty sandbox.

The sandbox receives the real `Workbook`, `Worksheet` and `Cell` objects wrapped
in `ClassInstance` policies, so sandbox code uses the `openpyxl` API it already
knows. What keeps it to one document is what the policies leave out: no
`load_workbook`, no `parent`, no styles or images, and a `save()` that takes no
filename.

The limits below are illustrative. `openpyxl` runs on the host, outside the
sandbox's `max_memory`, so a production host must choose limits for its own
workload, and consider what the loaded file itself may contain.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import KW_ONLY, dataclass
from inspect import signature
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import column_index_from_string
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from pydantic_monty import ClassInstance

__all__ = 'MAX_CELLS', 'MAX_COLUMNS', 'MAX_ROWS', 'MAX_SHEETS', 'DocumentWrapper', 'open_document'

MAX_ROWS = 10_000
"""Highest row index sandbox code may address; `openpyxl` allocates a cell for
every coordinate it is asked about, so the sheet's own 1,048,576 is too many."""
MAX_COLUMNS = 200
"""Highest column index sandbox code may address."""
MAX_CELLS = 100_000
"""Most cells one `iter_rows` or `iter_cols` call may materialize on the host."""
MAX_SHEETS = 20
"""Most worksheets sandbox code may grow the document to."""

ROW_ARGS = frozenset({'row', 'min_row', 'max_row'})
COLUMN_ARGS = frozenset({'column', 'min_col', 'max_col'})
WORKSHEET_METHODS = frozenset({'cell', 'append', 'iter_rows', 'iter_cols'})


def open_document(path: Path) -> DocumentWrapper:
    """Loads the workbook at `path` and wraps it for the sandbox; `path` is the
    only file the session can read, and the only one `save()` writes."""
    return DocumentWrapper(load_workbook(path), path=path)


@dataclass
class SpreadsheetWrapper(ClassInstance):
    """Base policy for every `openpyxl` object: wraps the worksheets and cells a
    call returns, and turns generators into lists so they can cross."""

    def convert_value(self, /, name: str, value: Any) -> Any:
        if isinstance(value, Worksheet):
            return WorksheetWrapper(value)
        elif isinstance(value, (Cell, MergedCell)):
            return cell_wrapper(value)
        elif isinstance(value, (list, tuple, Iterator)):
            return [self.convert_value(name, item) for item in value]  # pyright: ignore[reportUnknownVariableType]
        else:
            return value


@dataclass
class DocumentWrapper(SpreadsheetWrapper):
    """The workbook, and the root of everything the sandbox can reach."""

    value: Workbook
    _: KW_ONLY
    path: Path
    """Where the workbook was loaded from, the only place `save()` writes."""

    def __post_init__(self) -> None:
        self.lazy_attrs = {'sheetnames', 'worksheets', 'active'}
        self.allowed_methods = {'create_sheet', 'remove', 'save'}
        super().__post_init__()

    def call_method(self, name: str, args: tuple[Any, ...], kwargs: dict[str, Any]) -> Any:
        """`save()` writes back to `path`: `Workbook.save(filename)` is never
        reached with a filename from the sandbox."""
        if name == 'save':
            if args or kwargs:
                raise TypeError('save() takes no arguments, the document is saved in place')
            self.value.save(self.path)
        else:
            if name == 'create_sheet' and len(self.value.worksheets) >= MAX_SHEETS:
                raise ValueError(f'the document is limited to {MAX_SHEETS} worksheets')
            return super().call_method(name, args, kwargs)


@dataclass
class WorksheetWrapper(SpreadsheetWrapper):
    """One sheet. Its attributes are lazy, so they follow the sandbox's writes."""

    value: Worksheet

    def __post_init__(self) -> None:
        # `values` is left out: it walks the whole sheet with no bounds to check
        self.lazy_attrs = {'title', 'dimensions', 'min_row', 'max_row', 'min_column', 'max_column'}
        self.allowed_methods = WORKSHEET_METHODS
        super().__post_init__()

    def call_method(self, name: str, args: tuple[Any, ...], kwargs: dict[str, Any]) -> Any:
        """Checks every row and column index against the limits before calling."""
        if name in WORKSHEET_METHODS:
            self.check_bounds(name, args, kwargs)
        return super().call_method(name, args, kwargs)

    def check_bounds(self, name: str, args: tuple[Any, ...], kwargs: dict[str, Any]) -> None:
        """Raises `ValueError` for a call that would grow the sheet past
        `MAX_ROWS` x `MAX_COLUMNS`, or read more than `MAX_CELLS` at once.
        Arguments are bound to the method's own signature, so a positional
        index is checked like a keyword one."""
        arguments = signature(getattr(self.value, name)).bind(*args, **kwargs).arguments
        for arg_name, arg_value in arguments.items():
            limit = MAX_ROWS if arg_name in ROW_ARGS else MAX_COLUMNS if arg_name in COLUMN_ARGS else None
            if limit is not None and arg_value is not None:
                if not isinstance(arg_value, int) or arg_value > limit:
                    raise ValueError(f'{name}() {arg_name}={arg_value!r} is not an integer up to {limit}')
        if name == 'append':
            self.check_append(arguments['iterable'])
        elif name in {'iter_rows', 'iter_cols'}:
            # an omitted bound defaults to the sheet's own extent, which the loaded file decides
            rows = (arguments.get('max_row') or self.value.max_row) - (arguments.get('min_row') or 1) + 1
            columns = (arguments.get('max_col') or self.value.max_column) - (arguments.get('min_col') or 1) + 1
            if rows * columns > MAX_CELLS:
                raise ValueError(f'{name}() would read {rows} rows x {columns} columns, more than {MAX_CELLS} cells')

    def check_append(self, row: Any) -> None:
        """A dict row addresses columns by index or letter, so its keys are the
        columns it would allocate; a sequence row fills columns from the first."""
        columns: list[Any] = [len(row)]
        if isinstance(row, dict):
            columns = [column_index_from_string(key) if isinstance(key, str) else key for key in row]  # pyright: ignore[reportUnknownVariableType]
        if self.value.max_row >= MAX_ROWS or any(not isinstance(c, int) or c > MAX_COLUMNS for c in columns):
            raise ValueError(f'append() would grow the sheet beyond {MAX_ROWS} rows x {MAX_COLUMNS} columns')


def cell_wrapper(cell: Cell | MergedCell) -> ClassInstance:
    """A cell crosses as a snapshot of its position and value. Sandbox code
    writes with `ws.cell(row, column, value)`: setting `cell.value` changes the
    sandbox's copy only."""
    attrs = ('coordinate', 'row', 'column', 'value')
    if isinstance(cell, Cell):
        attrs += 'column_letter', 'data_type', 'number_format', 'is_date'
    return ClassInstance(cell, eager_attrs=attrs)
